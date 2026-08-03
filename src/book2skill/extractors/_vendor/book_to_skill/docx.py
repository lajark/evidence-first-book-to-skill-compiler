# SPDX-License-Identifier: MIT
#
# Selective port from virgiliojr94/book-to-skill
#   Original file: book_to_skill/parsers/docx.py
#   Repository: https://github.com/virgiliojr94/book-to-skill
#   Commit: 92b248fa5e7039d770d56630444310e36ff014e0
#   License: MIT (see LICENSES/MIT-upstream-virgilio-book-to-skill.txt)
#   Provenance ID: upstream-virgilio-book-to-skill
#
# Local modifications:
#   - Added provenance header.
#   - Adjusted import to use the vendored `exceptions` module within this package.
#   - Removed the `extract_docx` orchestrator and its `print` debug logging;
#     backend selection is handled by the DocxExtractor wrapper instead.
#   - XML safety scan uses one encoding decision per member and enforces
#     per-member / aggregate decompression budgets before reading XML bytes.

"""DOCX text extraction with optional python-docx backend and stdlib fallback.

Includes ``validate_docx_xml_safety`` to reject DTD/entity declarations
(defense against Billion Laughs / XXE in DOCX archives).
"""

from __future__ import annotations

import sys
import zipfile

from book2skill.extractors._vendor.book_to_skill.exceptions import ExtractionError

_XML_SUFFIXES = (".xml", ".rels")
_MAX_XML_MEMBER_BYTES = 16 * 1024 * 1024
_MAX_TOTAL_XML_BYTES = 64 * 1024 * 1024


def extract_docx_with_python_docx(docx_path: str) -> str | None:
    try:
        import docx

        document = docx.Document(docx_path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))
        return "\n".join(parts)
    except ImportError:
        return None
    except Exception as e:
        print(
            f"  [warn] extract_docx_with_python_docx failed: {type(e).__name__}: {e}",
            file=sys.stderr,
        )
        return None


def extract_docx_with_zipfile(docx_path: str) -> str | None:
    try:
        import xml.etree.ElementTree as ET

        with zipfile.ZipFile(docx_path) as zf:
            xml_bytes = zf.read("word/document.xml")
        root = ET.fromstring(xml_bytes)
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        parts: list[str] = []

        def emit_block(elem) -> None:
            # Walk block content in document order. Paragraphs join their runs;
            # tables emit one tab-joined line per row (same row format as the
            # python-docx path, but order-preserving — python-docx appends all
            # tables last). Unknown wrappers (e.g. <w:sdt> content controls) are
            # recursed into so their paragraphs/tables are not lost; <w:p> and
            # <w:tbl> are NOT recursed into, so table-cell paragraphs are not
            # double-counted. Cell text concatenates the cell's runs; nested
            # tables fold into the parent cell and are also emitted standalone
            # (rare; best-effort).
            for child in elem:
                tag = child.tag
                if tag == f"{ns}p":
                    texts = [t.text for t in child.iter(f"{ns}t") if t.text]
                    if texts:
                        parts.append("".join(texts))
                elif tag == f"{ns}tbl":
                    for row in child.iter(f"{ns}tr"):
                        cells = []
                        for cell in row.iter(f"{ns}tc"):
                            cell_texts = [
                                t.text for t in cell.iter(f"{ns}t") if t.text
                            ]
                            cells.append("".join(cell_texts).strip())
                        if any(cells):
                            parts.append("\t".join(cells))
                else:
                    emit_block(child)

        body = root.find(f"{ns}body")
        emit_block(body if body is not None else root)
        return "\n".join(parts) if parts else None
    except Exception as e:
        print(
            f"  [warn] extract_docx_with_zipfile failed: {type(e).__name__}: {e}",
            file=sys.stderr,
        )
        return None


def validate_docx_xml_safety(docx_path: str) -> None:
    """Scan all XML files in the DOCX zip archive to prevent XML Entity
    Expansion (Billion Laughs) and XXE injections."""
    try:
        with zipfile.ZipFile(docx_path) as zf:
            total_xml_bytes = 0
            for info in zf.infolist():
                if not info.filename.endswith(_XML_SUFFIXES):
                    continue
                total_xml_bytes += info.file_size
                if (
                    info.file_size > _MAX_XML_MEMBER_BYTES
                    or total_xml_bytes > _MAX_TOTAL_XML_BYTES
                ):
                    raise ExtractionError(
                        f"Security validation failed: XML file '{info.filename}' "
                        "exceeds the XML safety budget."
                    )
                content = _decode_xml_for_safety(zf.read(info)).casefold()
                if "<!doctype" in content or "<!entity" in content:
                    raise ExtractionError(
                        f"Security validation failed: XML file '{info.filename}' "
                        "in DOCX archive contains forbidden DTD or entity declarations."
                    )
    except zipfile.BadZipFile as e:
        raise ExtractionError(f"Invalid DOCX file: {e}")
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Error during security validation of DOCX archive: {e}")


def _decode_xml_for_safety(xml_bytes: bytes) -> str:
    """Decode an XML member once using BOM or byte-order inspection.

    XML declarations are ASCII-compatible, so examining the first four bytes
    covers BOM-less UTF-16/32 documents without attempting several full-size
    decodes of the same archive member.
    """
    if xml_bytes.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        encoding = "utf-32"
    elif xml_bytes.startswith((b"\xff\xfe", b"\xfe\xff")):
        encoding = "utf-16"
    elif xml_bytes.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    elif xml_bytes.startswith(b"<\x00?\x00"):
        encoding = "utf-16le"
    elif xml_bytes.startswith(b"\x00<\x00?"):
        encoding = "utf-16be"
    elif xml_bytes.startswith(b"<\x00\x00\x00"):
        encoding = "utf-32le"
    elif xml_bytes.startswith(b"\x00\x00\x00<"):
        encoding = "utf-32be"
    else:
        encoding = "utf-8"
    return xml_bytes.decode(encoding, errors="ignore")
